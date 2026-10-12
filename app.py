import csv
import hashlib
import io
import re
import secrets
import subprocess
import time
import unicodedata
from contextlib import closing
from datetime import datetime
from pathlib import Path

import pypdfium2 as pdfium
from flask import Flask, abort, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.serving import WSGIRequestHandler, make_server

ROOT = Path(__file__).resolve().parent
FIELDS = ["書類種別", "発行日", "書類番号", "取引先名", "注文番号", "合計金額"]
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024
workspaces = {}
with (ROOT / "samples/ledger.csv").open(encoding="utf-8", newline="") as ledger_file:
    LEDGER = list(csv.DictReader(ledger_file))


def normalize(value):
    return unicodedata.normalize("NFKC", value).strip()


def parse_fields(pages):
    values = dict.fromkeys(FIELDS, "")
    evidence = {}
    for number, page in enumerate(pages, 1):
        for line in page["text"].splitlines():
            compact = re.sub(r"\s+", "", normalize(line))
            for field in FIELDS:
                match = re.fullmatch(re.escape(field) + r"[:：](.+)", compact)
                if match and not values[field]:
                    values[field] = match[1][:200]
                    evidence[field] = {"page": number, "text": line}
    return values, evidence


def read_pdf(data, name):
    pages = []
    try:
        with closing(pdfium.PdfDocument(data)) as pdf:
            if not 1 <= len(pdf) <= 5:
                raise ValueError("PDFは1〜5ページにしてください。")
            for page in pdf:
                with closing(page):
                    width, height = page.get_size()
                    if width <= 0 or height <= 0 or max(width, height) > 1200:
                        raise ValueError("ページサイズはA3相当までにしてください。")
                    with closing(page.get_textpage()) as textpage:
                        text = textpage.get_text_bounded()
                    with closing(page.render(scale=2.5)) as bitmap:
                        image = bitmap.to_pil()
                        png = io.BytesIO()
                        image.save(png, format="PNG")
                    method = "テキスト抽出"
                    if not text.strip():
                        method = "日本語OCR"
                        try:
                            result = subprocess.run(
                                ["tesseract", "stdin", "stdout", "-l", "jpn+eng", "--psm", "6"],
                                input=png.getvalue(),
                                capture_output=True,
                                timeout=45,
                                check=True,
                            )
                            text = result.stdout.decode("utf-8")
                        except (OSError, subprocess.SubprocessError) as exc:
                            raise ValueError(
                                "OCRを利用できません。Tesseractとjpn/engを確認してください。"
                            ) from exc
                    pages.append({"text": text, "method": method, "png": png.getvalue()})
    except pdfium.PdfiumError as exc:
        raise ValueError("PDFを開けません。暗号化や破損のないPDFを選んでください。") from exc
    values, evidence = parse_fields(pages)
    return {
        "id": hashlib.sha256(data).hexdigest(),
        "name": name[:150],
        "pages": pages,
        "values": values,
        "evidence": evidence,
        "confirmed": False,
        "excluded": False,
        "reason": "",
    }


def reconciliation(doc, docs):
    values = doc["values"]
    errors = [f"{field}が未取得です" for field in FIELDS if not values[field]]
    date = amount = None
    if values["書類種別"] not in ("請求書", "納品書"):
        errors.append("書類種別は請求書または納品書にしてください")
    try:
        raw_date = normalize(values["発行日"])
        if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", raw_date):
            raise ValueError("date format")
        date = datetime.strptime(raw_date, "%Y-%m-%d")
    except ValueError:
        errors.append("発行日は有効なYYYY-MM-DDで入力してください")
    raw_amount = normalize(values["合計金額"])
    if re.fullmatch(r"(?:[0-9]+|[0-9]{1,3}(?:,[0-9]{3})+)(?:円)?", raw_amount):
        amount = int(raw_amount.removesuffix("円").replace(",", ""))
        if amount > 999999999999:
            errors.append("金額は12桁以内にしてください")
    else:
        errors.append("金額は0以上の整数（円）で入力してください")
    key = tuple(normalize(values[f]) for f in ("書類種別", "取引先名", "書類番号"))
    duplicate = (
        not doc["excluded"]
        and all(key)
        and any(
            other["id"] != doc["id"]
            and not other["excluded"]
            and tuple(normalize(other["values"][f]) for f in ("書類種別", "取引先名", "書類番号")) == key
            for other in docs
        )
    )
    if duplicate:
        errors.append("重複候補です。不要な書類を除外してください")
    matches = [
        row
        for row in LEDGER
        if normalize(row["取引先名"]) == normalize(values["取引先名"])
        and normalize(row["注文番号"]) == normalize(values["注文番号"])
    ]
    state = "未照合"
    expected = None
    if len(matches) == 1 and amount is not None:
        expected = int(matches[0]["予定金額"])
        state = "一致" if amount == expected else "金額不一致"
    return {"errors": errors, "state": state, "expected": expected, "date": date, "amount": amount}


def snapshot(workspace):
    docs = list(workspace["docs"].values())
    output = []
    for doc in docs:
        result = reconciliation(doc, docs)
        if result["errors"]:
            doc["confirmed"] = False
        output.append(
            {k: v for k, v in doc.items() if k != "pages"}
            | {
                "pages": [{"text": p["text"], "method": p["method"]} for p in doc["pages"]],
                "result": {k: v for k, v in result.items() if k not in ("date", "amount")},
            }
        )
    return output


@app.before_request
def guard():
    if request.host not in ("127.0.0.1:5055", "localhost:5055"):
        abort(403)
    now = time.monotonic()
    for token in list(workspaces):
        if now - workspaces[token]["touched"] > 1800:
            del workspaces[token]
    if request.path.startswith("/api/"):
        token = request.headers.get("X-Workspace")
        if request.path == "/api/export" and request.method == "POST":
            token = token or request.form.get("workspace")
        if token not in workspaces:
            abort(410, description="作業が期限切れです。画面を再読み込みしてください。")
        workspaces[token]["touched"] = now


@app.after_request
def headers(response):
    response.headers["Cache-Control"] = "no-store"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; img-src 'self' blob:; script-src 'self'; style-src 'self'; "
        "connect-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'"
    )
    return response


@app.errorhandler(413)
def too_large(_error):
    return jsonify(error="PDFは10MB以内にしてください。"), 413


@app.get("/")
def index():
    if len(workspaces) >= 8:
        abort(503, description="作業数の上限です。30分後に再試行してください。")
    token = secrets.token_urlsafe(32)
    workspaces[token] = {"docs": {}, "touched": time.monotonic()}
    return render_template("index.html", token=token, fields=FIELDS)


def workspace():
    token = request.headers.get("X-Workspace") or request.form.get("workspace")
    return workspaces[token]


def add_pdf(data, name):
    current = workspace()
    digest = hashlib.sha256(data).hexdigest()
    if digest in current["docs"]:
        return False
    if len(current["docs"]) >= 12:
        raise ValueError("1回の作業は12書類までです。作業を終了してやり直してください。")
    current["docs"][digest] = read_pdf(data, name)
    return True


@app.post("/api/load")
def load():
    try:
        if request.form.get("sample") == "yes":
            paths = sorted((ROOT / "samples").glob("*.pdf"))
            count = sum(add_pdf(path.read_bytes(), path.name) for path in paths)
        else:
            file = request.files.get("pdf")
            if not file or not file.filename.lower().endswith(".pdf"):
                raise ValueError("PDFファイルを選んでください。")
            name = file.filename.replace("\\", "/").rsplit("/", 1)[-1]
            count = int(add_pdf(file.read(), name))
        return jsonify(
            docs=snapshot(workspace()), message=f"{count}書類を追加しました。同じPDFは追加しません。"
        )
    except ValueError as exc:
        return jsonify(error=str(exc)), 400


@app.get("/api/image/<doc_id>/<int:page>")
def page_image(doc_id, page):
    doc = workspace()["docs"].get(doc_id)
    if doc is None or not 1 <= page <= len(doc["pages"]):
        abort(404)
    return send_file(io.BytesIO(doc["pages"][page - 1]["png"]), mimetype="image/png")


@app.post("/api/document/<doc_id>")
def edit(doc_id):
    current = workspace()
    doc = current["docs"].get(doc_id)
    if doc is None:
        abort(404)
    data = request.get_json()
    action = data.get("action")
    if action == "save":
        doc["values"] = {field: str(data["values"].get(field, "")).strip()[:200] for field in FIELDS}
        doc["reason"] = str(data.get("reason", "")).strip()[:1000]
        doc["confirmed"] = False
    elif action == "exclude":
        doc["excluded"] = not doc["excluded"]
        doc["confirmed"] = False
    elif action == "confirm":
        result = reconciliation(doc, list(current["docs"].values()))
        if doc["excluded"] or result["errors"] or (result["state"] != "一致" and not doc["reason"]):
            return jsonify(error="項目・重複を修正してください。不一致・未照合には確認理由が必要です。"), 400
        doc["confirmed"] = True
    else:
        abort(400)
    return jsonify(docs=snapshot(current))


@app.post("/api/export")
def export():
    current = workspace()
    snapshot(current)
    docs = list(current["docs"].values())
    wb = Workbook()
    sheet = wb.active
    sheet.title = "確認済み書類"
    sheet.append([*FIELDS, "元ファイル名", "根拠ページ", "照合状態", "確認理由"])
    for doc in docs:
        result = reconciliation(doc, docs)
        if not doc["confirmed"] or doc["excluded"] or result["errors"]:
            continue
        if result["state"] != "一致" and not doc["reason"]:
            continue
        values = doc["values"]
        row = [
            values["書類種別"],
            result["date"],
            values["書類番号"],
            values["取引先名"],
            values["注文番号"],
            result["amount"],
            doc["name"],
            ", ".join(f"{f}:p{e['page']}" for f, e in doc["evidence"].items()),
            result["state"],
            doc["reason"],
        ]
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            # 元書類・人の入力を数式として評価させない。
            if isinstance(cell.value, str):
                cell.data_type = "s"
        sheet.cell(sheet.max_row, 2).number_format = "yyyy-mm-dd"
        sheet.cell(sheet.max_row, 6).number_format = '#,##0"円"'
    if sheet.max_row == 1:
        return jsonify(error="出力できる確認済み書類がありません。"), 400
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=11, bold=cell.row == 1)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="DCEFE9")
    for column, width in zip("ABCDEFGHIJ", [16, 16, 22, 24, 22, 18, 35, 65, 18, 55], strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, download_name="confirmed-documents.xlsx", as_attachment=True)


@app.post("/api/clear")
def clear():
    workspace()["docs"].clear()
    return jsonify(docs=[])


class QuietHandler(WSGIRequestHandler):
    def log(self, _type, _message, *args):
        pass


if __name__ == "__main__":
    # localhost専用。本文・ファイル名をアクセスログへ残さない。
    with make_server("127.0.0.1", 5055, app, request_handler=QuietHandler) as server:
        print("書類照合デモ: http://127.0.0.1:5055", flush=True)
        server.serve_forever()
